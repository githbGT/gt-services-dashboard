"""
Step 1: read the service catalog (Excel) and write a clean CSV.

The Excel file is opened read-only and never modified.
Every row with a service name is read, so new services added anywhere in the
sheet (inside or below the table) are picked up automatically.
"""
import csv
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import config  # noqa: E402

CSV_FIELDS = ["name", "url", "unit", "description", "type", "audience", "online",
              "ch_branch", "ch_phone", "ch_email", "ch_mail", "ch_kiosk",
              "level", "level_name", "usage_total", "usage_digital",
              "usage_frontal", "usage_phone"]


def norm(v):
    """Clean a cell value: None -> '', collapse whitespace incl. non-breaking spaces."""
    if v is None:
        return ""
    s = str(v).replace("\xa0", " ").replace("*", " ")
    return re.sub(r"\s+", " ", s).strip()


def find_header_row(rows):
    key = norm(config.COLUMNS["name"])
    for i, row in enumerate(rows[:5]):
        if any(norm(v).startswith(key) for v in row):
            return i
    sys.exit(f"שגיאה: לא נמצאה שורת כותרות עם '{key}' בחמש השורות הראשונות.")


def map_columns(header):
    """Map each field in config.COLUMNS to a column index in the header row."""
    headers = [norm(h) for h in header]
    found = {}
    for field, text in config.COLUMNS.items():
        text = norm(text)
        exact = [i for i, h in enumerate(headers) if h == text]
        prefix = [i for i, h in enumerate(headers) if h.startswith(text)]
        if not (exact or prefix):
            sys.exit(f"שגיאה: לא נמצאה עמודה '{text}'. אם שם העמודה השתנה, "
                     f"יש לעדכן אותו ב-config.py תחת COLUMNS.")
        found[field] = (exact or prefix)[0]
    return found


def to_int(v, warn, name, field):
    s = norm(v).replace(",", "")
    if s == "":
        return 0
    try:
        return int(round(float(s)))
    except ValueError:
        warn(f"'{name}': ערך לא מספרי בעמודה {field}: '{v}' (נספר כ-0)")
        return 0


def clean_url(v):
    """Return the first http(s) URL in the cell, or '' ("אין", phone numbers etc.)."""
    m = re.search(r"https?://[^\s\\\"'<>]+", str(v or ""))
    return m.group(0).rstrip(".,;)") if m else ""


def extract(path):
    warnings, new_values = [], {}

    def warn(msg):
        warnings.append(msg)

    def new_value(kind, val, name):
        new_values.setdefault((kind, val), []).append(name)

    # read_only + data_only: values only, the workbook is never written
    wb = load_workbook(path, read_only=True, data_only=True)
    if config.CATALOG_SHEET not in wb.sheetnames:
        sys.exit(f"שגיאה: לא נמצא הגיליון '{config.CATALOG_SHEET}'. "
                 f"גיליונות בקובץ: {', '.join(wb.sheetnames)}")
    ws = wb[config.CATALOG_SHEET]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()

    header_row = find_header_row(rows)
    col = map_columns(rows[header_row])

    records = []
    for row in rows[header_row + 1:]:
        def get(field, row=row):
            i = col[field]
            return row[i] if i < len(row) else None

        name = norm(get("name"))
        if not name:
            continue

        audience_raw = norm(get("audience"))
        audience = config.AUDIENCE_MAP.get(audience_raw, audience_raw)
        if audience_raw and audience_raw not in config.AUDIENCE_MAP:
            new_value("קהל יעד", audience_raw, name)
        if not audience:
            warn(f"'{name}': חסר קהל יעד")

        stype = norm(get("type"))
        if stype and stype not in config.KNOWN_TYPES:
            new_value("סוג שירות", stype, name)

        m = re.match(r"\s*([1-5])", norm(get("level")))
        level = int(m.group(1)) if m else ""
        if not m:
            warn(f"'{name}': רמת דיגיטליות לא מזוהה: '{norm(get('level'))}' "
                 f"(השירות לא ייכלל בסולם הדיגיטליות)")

        online = norm(get("online"))
        if online not in ("כן", "לא"):
            warn(f"'{name}': ערך 'מקוון' חסר או לא תקין: '{online}'")
            online = ""

        channels = {}
        for f in ("ch_branch", "ch_phone", "ch_email", "ch_mail", "ch_kiosk"):
            v = norm(get(f))
            if v not in ("כן", "לא", ""):
                warn(f"'{name}': ערך לא תקין בעמודה {config.COLUMNS[f]}: '{v}'")
            channels[f] = v if v in ("כן", "לא") else ""

        rec = {
            "name": name,
            "url": clean_url(get("url")),
            "unit": norm(get("unit")),
            "description": norm(get("description")),
            "type": stype,
            "audience": audience,
            "online": online,
            **channels,
            "level": level,
            "level_name": config.LEVEL_NAMES.get(level, ""),
        }
        for f in ("usage_total", "usage_digital", "usage_frontal", "usage_phone"):
            rec[f] = to_int(get(f), warn, name, f)

        parts = rec["usage_digital"] + rec["usage_frontal"] + rec["usage_phone"]
        if parts and parts != rec["usage_total"]:
            warn(f"'{name}': סה\"כ שימוש {rec['usage_total']:,} אבל סכום הערוצים {parts:,}")
        records.append(rec)

    names = [r["name"] for r in records]
    for n in sorted({n for n in names if names.count(n) > 1}):
        warn(f"שם שירות כפול: '{n}'")

    return records, warnings, new_values


def write_csv(records, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    # utf-8-sig so the file also opens correctly in Excel; \n line endings for clean git diffs
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=CSV_FIELDS, lineterminator="\n")
        w.writeheader()
        w.writerows(records)


def main(src=None):
    src = Path(src) if src else config.CATALOG_PATH
    if not src.exists():
        sys.exit(f"שגיאה: קובץ הקטלוג לא נמצא: {src}\n"
                 f"אפשר להעביר נתיב: python update.py \"<נתיב לקובץ>\"")
    records, warnings, new_values = extract(src)
    write_csv(records, config.PUBLIC_CSV)

    print(f"שלב 1: נקראו {len(records)} שירותים -> {config.PUBLIC_CSV.relative_to(config.ROOT)}")
    if new_values:
        print("\nערכים חדשים (יופיעו בדשבורד כפי שהם; אפשר להוסיף מיפוי ב-config.py):")
        for (kind, val), who in new_values.items():
            print(f"  {kind}: '{val}' ({len(who)} שירותים)")
    if warnings:
        print(f"\nהערות על הנתונים ({len(warnings)}):")
        for w in warnings:
            print("  - " + w)
    return records


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else None)
